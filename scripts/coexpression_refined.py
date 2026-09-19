# -*- coding: utf-8 -*-
"""
共表达基因的进一步筛选 (不依赖 top20):
  在全部背景基因里, 保留与至少 MIN_MARKERS 个 marker 显著正相关的基因
  (Pearson r > R_THRESHOLD; n = 6 时 r > 0.811 对应双侧 P < 0.05),
  然后按相关谱聚类分成两个模块 (模块 A / 模块 B) 画热图, 并导出基因表。

输出 (在 <WORK_ROOT>\\当天日期\\):
  Fig_coexpression_refined.pdf / .svg / .png
  Fig_coexpression_refined_genes.csv      每个入选基因与各 marker 的 r、P、BH q, 以及模块归属
  Fig_coexpression_refined.json           供 Word 文档生成脚本读取
"""

import os
import sys
import json
import datetime
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm
from matplotlib.patches import Rectangle

# ---------------- 工作目录 ----------------
WORK_ROOT = r"C:\Users\23027\Desktop\11\画各种图"
os.makedirs(WORK_ROOT, exist_ok=True)
os.chdir(WORK_ROOT)
DATA_DIR = os.path.join(WORK_ROOT, "data")
os.makedirs(DATA_DIR, exist_ok=True)
OUT_DIR = os.path.join(WORK_ROOT, datetime.date.today().strftime("%Y-%m-%d"))
os.makedirs(OUT_DIR, exist_ok=True)


def _pick(stem):
    for ext in (".txt", ".tsv", ".csv", ".xlsx"):
        cand = os.path.join(DATA_DIR, stem + ext)
        if os.path.isfile(cand):
            return cand
    return os.path.join(DATA_DIR, stem + ".xlsx")


EXPR_FILE = _pick("expression_matrix")
MARKER_FILE = _pick("marker_genes")
if len(sys.argv) >= 3:
    EXPR_FILE, MARKER_FILE = sys.argv[1], sys.argv[2]
OUT_STEM = os.path.join(OUT_DIR, "Fig_coexpression_refined")

# ---------------- 筛选参数 ----------------
R_THRESHOLD = 0.811    # n = 6 时 |r| > 0.811 <=> 双侧 P < 0.05
MIN_MARKERS = 3        # 至少与几个 marker 显著正相关; 改成 4 更严
HIGHLIGHT = ["Chr07Ag005522", "Chr06Bg002494"]   # 重点候选, 热图上用红色标出
GENE_LABEL_SIZE = 7

# 两个模块的定义 (用于给聚类结果命名, 不参与筛选)
MODULE_A_MARKERS = ["Chr02Bg000869", "Chr02Bg005918", "Chr02Bg006322", "Chr06Ag003232"]
MODULE_B_MARKERS = ["Chr03Ag002790"]

# ---------------- 字体 ----------------
plt.rcParams["font.family"] = "serif"
plt.rcParams["font.serif"] = ["Times New Roman", "SimSun", "Liberation Serif"]
plt.rcParams["font.sans-serif"] = ["SimSun"]
plt.rcParams["mathtext.fontset"] = "stix"
plt.rcParams["font.weight"] = "bold"
plt.rcParams["axes.labelweight"] = "bold"
plt.rcParams["axes.titleweight"] = "bold"
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["font.size"] = 10.5
plt.rcParams["pdf.fonttype"] = 42
plt.rcParams["ps.fonttype"] = 42
plt.rcParams["svg.fonttype"] = "none"


# ---------------- 读取 ----------------
def _rows_from_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xltx"):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    with open(path, "r", encoding="utf-8-sig") as fh:
        lines = [ln.rstrip("\r\n") for ln in fh if ln.strip()]
    sep = "\t" if "\t" in lines[0] else ("," if "," in lines[0] else None)
    return [ln.split(sep) if sep else [ln] for ln in lines]


def read_matrix(path):
    rows = _rows_from_file(path)
    header = [h for h in rows[0] if h not in (None, "")]
    samples = [str(h) for h in header[1:]]
    genes, values = [], []
    for r in rows[1:]:
        if r[0] in (None, ""):
            continue
        genes.append(str(r[0]).strip())
        values.append([float(r[i + 1]) for i in range(len(samples))])
    return samples, genes, np.array(values, dtype=float)


def read_marker_ids(path):
    rows = _rows_from_file(path)
    ids = []
    for i, r in enumerate(rows):
        first = str(r[0]).strip() if r and r[0] is not None else ""
        if not first:
            continue
        if i == 0 and (first.startswith("#") or first.lower() in ("id", "gene", "geneid")):
            continue
        ids.append(first)
    return ids


def pearson_vs_matrix(v, M):
    vc = v - v.mean()
    Mc = M - M.mean(axis=1, keepdims=True)
    denom = np.sqrt((Mc ** 2).sum(axis=1)) * np.sqrt((vc ** 2).sum())
    denom[denom == 0] = np.nan
    return (Mc @ vc) / denom


samples, genes, counts = read_matrix(EXPR_FILE)
marker_ids = read_marker_ids(MARKER_FILE)
cpm = counts / counts.sum(axis=0, keepdims=True) * 1e6
logcpm = np.log2(cpm + 1.0)
gene_index = {g: i for i, g in enumerate(genes)}
missing = [g for g in marker_ids if g not in gene_index]
if missing:
    raise ValueError("marker 基因不在表达矩阵中: {}".format(missing))
other_genes = [g for g in genes if g not in set(marker_ids)]
other_mat = logcpm[[gene_index[g] for g in other_genes]]

# ---------------- 相关系数、P 值、BH q ----------------
R = np.vstack([pearson_vs_matrix(logcpm[gene_index[m]], other_mat) for m in marker_ids])
n_sample = logcpm.shape[1]
df_t = n_sample - 2
with np.errstate(divide="ignore", invalid="ignore"):
    tstat = R * np.sqrt(df_t) / np.sqrt(np.maximum(1.0 - R ** 2, 1e-12))
from scipy import stats as _stats
P = 2.0 * _stats.t.sf(np.abs(tstat), df_t)
flat = P.ravel()
order = np.argsort(flat)
mtest = flat.size
Q = np.empty(mtest)
running = 1.0
for rank, pos in enumerate(order[::-1]):
    running = min(running, flat[pos] * mtest / (mtest - rank))
    Q[pos] = running
Q = Q.reshape(P.shape)

# ---------------- 筛选: 与 >= MIN_MARKERS 个 marker 显著正相关 ----------------
n_sig = (R > R_THRESHOLD).sum(axis=0)
keep = np.where(n_sig >= MIN_MARKERS)[0]
sel_names = [other_genes[j] for j in keep]
Rsel = R[:, keep]
n_row, n_col = Rsel.shape
if n_col == 0:
    raise SystemExit("没有基因满足筛选条件, 请放宽 MIN_MARKERS 或 R_THRESHOLD")

# ---------------- 聚类排序 + 模块划分 ----------------
from scipy.cluster.hierarchy import linkage, leaves_list, fcluster
from scipy.spatial.distance import pdist
lk = linkage(np.nan_to_num(pdist(Rsel.T, metric="correlation"), nan=1.0),
             method="average", optimal_ordering=True)
col_order = leaves_list(lk)
clusters = fcluster(lk, 2, criterion="maxclust") if n_col >= 2 else np.ones(n_col, int)

# 用模块定义给两个聚类命名: 与模块 A 的 4 个 marker 平均 r 更高的叫 A, 另一个叫 B
mi = {m: i for i, m in enumerate(marker_ids)}
idx_a = [mi[m] for m in MODULE_A_MARKERS if m in mi]
idx_b = [mi[m] for m in MODULE_B_MARKERS if m in mi]
score = {}
for c in np.unique(clusters):
    sub = Rsel[:, clusters == c]
    score[c] = sub[idx_a].mean() - (sub[idx_b].mean() if idx_b else 0.0)
c_sorted = sorted(score, key=lambda c: -score[c])
module_name = {c_sorted[0]: "A"}
if len(c_sorted) > 1:
    module_name[c_sorted[1]] = "B"
modules = [module_name[c] for c in clusters]

Rsel = Rsel[:, col_order]
sel_names = [sel_names[j] for j in col_order]
modules = [modules[j] for j in col_order]
keep = keep[col_order]

# ---------------- 热图 ----------------
cmap = LinearSegmentedColormap.from_list(
    "npg_blue_white_red", ["#3C5488", "#FFFFFF", "#E64B35"])
norm = TwoSlopeNorm(vmin=-1.0, vcenter=0.0, vmax=1.0)
MOD_COLOR = {"A": "#91ABD2", "B": "#FBD3A3"}

fig_w = max(5.5, n_col * 0.19 + 2.2)
fig_h = n_row * 0.30 + 2.5
fig = plt.figure(figsize=(fig_w, fig_h), dpi=300)
gs = fig.add_gridspec(2, 1, height_ratios=[0.08, 1.0], hspace=0.04,
                      left=0.17, right=0.86, top=0.90, bottom=0.30)
ax_mod = fig.add_subplot(gs[0, 0])
ax = fig.add_subplot(gs[1, 0], sharex=ax_mod)

im = ax.imshow(Rsel, cmap=cmap, norm=norm, aspect="auto", interpolation="nearest")
ax.set_xticks(np.arange(n_col))
ax.set_xticklabels(sel_names, rotation=90, fontsize=GENE_LABEL_SIZE, fontweight="bold")
for lab in ax.get_xticklabels():
    if lab.get_text() in HIGHLIGHT:
        lab.set_color("#E64B35")
ax.set_yticks(np.arange(n_row))
ax.set_yticklabels(marker_ids, fontsize=8, fontweight="bold")
ax.set_xticks(np.arange(-0.5, n_col, 1), minor=True)
ax.set_yticks(np.arange(-0.5, n_row, 1), minor=True)
ax.grid(which="minor", color="white", linewidth=0.4)
ax.tick_params(which="minor", length=0)
ax.tick_params(which="major", length=2.0, width=0.8)
for side in ("left", "bottom", "right", "top"):
    ax.spines[side].set_linewidth(0.8)
ax.set_xlabel("Genes positively correlated with >= {} markers (r > {}, n = {})".format(
    MIN_MARKERS, R_THRESHOLD, n_col), fontsize=10.5, labelpad=6, fontweight="bold")
ax.set_ylabel("Marker genes", fontsize=10.5, fontweight="bold")

# 模块注释条: 连续同模块的列合并成一段
ax_mod.set_xlim(-0.5, n_col - 0.5)
ax_mod.set_ylim(0, 1)
ax_mod.axis("off")
start = 0
for j in range(1, n_col + 1):
    if j == n_col or modules[j] != modules[start]:
        m = modules[start]
        ax_mod.add_patch(Rectangle((start - 0.5, 0), j - start, 1,
                                   facecolor=MOD_COLOR.get(m, "#CCCCCC"),
                                   edgecolor="white", linewidth=0.6))
        ax_mod.text((start + j - 1) / 2.0, 0.5, "Module {} (n = {})".format(m, j - start),
                    ha="center", va="center", fontsize=8, fontweight="bold")
        start = j

cbar = fig.colorbar(im, ax=ax, fraction=0.035, pad=0.02,
                    ticks=[-1.0, -0.5, 0.0, 0.5, 1.0])
cbar.set_label("Pearson r", fontsize=10.5, fontweight="bold")
cbar.ax.tick_params(labelsize=9)
for lab in cbar.ax.get_yticklabels():
    lab.set_fontweight("bold")
cbar.outline.set_linewidth(0.8)

# 输出文件不写入任何软件来源信息
_META = {"pdf": {"Creator": None, "Producer": None},
         "svg": {"Creator": None, "Date": None},
         "png": {"Software": None}}
for ext in ("pdf", "svg", "png"):
    path = "{}.{}".format(OUT_STEM, ext)
    fig.savefig(path, format=ext, bbox_inches="tight", pad_inches=0.05,
                metadata=_META[ext])
    print("saved:", path)
plt.close(fig)

# ---------------- 基因表 (CSV) ----------------
csv_path = OUT_STEM + "_genes.csv"
with open(csv_path, "w", encoding="utf-8-sig") as fh:
    head = ["gene", "module", "n_sig_markers", "highlight"]
    for m in marker_ids:
        head += ["r_" + m, "P_" + m, "BH_q_" + m]
    fh.write(",".join(head) + "\n")
    for k, j in enumerate(keep):
        g = other_genes[j]
        cells = [g, modules[k], str(int(n_sig[j])), "yes" if g in HIGHLIGHT else "no"]
        for i in range(n_row):
            cells += ["{:.4f}".format(R[i, j]), "{:.4g}".format(P[i, j]), "{:.4g}".format(Q[i, j])]
        fh.write(",".join(cells) + "\n")
print("saved:", csv_path)

# ---------------- JSON (供 Word 文档) ----------------
records = []
for k, j in enumerate(keep):
    g = other_genes[j]
    records.append({
        "gene": g, "module": modules[k], "n_sig": int(n_sig[j]),
        "highlight": g in HIGHLIGHT,
        "r": [round(float(R[i, j]), 3) for i in range(n_row)],
        "P": [float("{:.3g}".format(P[i, j])) for i in range(n_row)],
    })
mod_summary = {}
for m in ("A", "B"):
    cols = [k for k in range(n_col) if modules[k] == m]
    if cols:
        mod_summary[m] = {"n": len(cols),
                          "mean_r": [round(float(Rsel[i, cols].mean()), 2) for i in range(n_row)]}
json_path = OUT_STEM + ".json"
with open(json_path, "w", encoding="utf-8") as fh:
    json.dump({"r_threshold": R_THRESHOLD, "min_markers": MIN_MARKERS,
               "n_background": len(other_genes), "n_selected": n_col,
               "marker_order": list(marker_ids), "highlight": HIGHLIGHT,
               "modules": mod_summary, "genes": records}, fh, ensure_ascii=False, indent=1)
print("saved:", json_path)

# ---------------- 关键信息 ----------------
print("\nrule: r > {} with >= {} markers, from {} background genes".format(
    R_THRESHOLD, MIN_MARKERS, len(other_genes)))
print("selected genes:", n_col)
for m, d in mod_summary.items():
    print("Module {}: {} genes, mean r with markers: {}".format(m, d["n"], d["mean_r"]))
print("n_sig = 3:", int((n_sig[keep] == 3).sum()), " n_sig = 4:", int((n_sig[keep] == 4).sum()),
      " n_sig >= 5:", int((n_sig[keep] >= 5).sum()))
print("highlight genes:", [(g, modules[sel_names.index(g)]) for g in HIGHLIGHT if g in sel_names])
