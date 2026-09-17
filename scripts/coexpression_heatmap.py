# -*- coding: utf-8 -*-
"""
6 个 marker 基因与其余基因的共表达热图。

流程:
  1. 读入表达矩阵 (436 基因 x 6 样本) 与 marker 列表 (6 基因)
  2. CPM 标准化后 log2(CPM+1)  (N 组与 P 组文库大小相差约 4 倍, 必须先标准化)
  3. 每个 marker 与其余基因逐一计算 Pearson 相关系数
  4. 每个 marker 取 |r| 最大的前 TOPN 个基因, 取并集作为热图的列
  5. 对列做层次聚类排序, 输出矢量热图 (PDF + SVG) 与完整相关系数表 (CSV)

用法:
  python coexpression_heatmap.py [表达矩阵文件] [marker文件]   (txt 或 xlsx 均可)
  不带参数时从 <WORK_ROOT>\\data 读取, 结果写入 <WORK_ROOT>\\当天日期 文件夹。
"""

import os
import sys
import datetime
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm

# ---------------- 工作目录 ----------------
# 直接把工作目录切到这里, 之后所有相对路径都以它为准:
#   <WORK_ROOT>\data\          输入数据 (xlsx)
#   <WORK_ROOT>\YYYY-MM-DD\    当天的输出, 每天自动新建一个文件夹
WORK_ROOT = r"C:\Users\23027\Desktop\11\画各种图"
os.makedirs(WORK_ROOT, exist_ok=True)
os.chdir(WORK_ROOT)
DATA_DIR = os.path.join(WORK_ROOT, "data")
os.makedirs(DATA_DIR, exist_ok=True)
OUT_DIR = os.path.join(WORK_ROOT, datetime.date.today().strftime("%Y-%m-%d"))
os.makedirs(OUT_DIR, exist_ok=True)

# 输入文件: 同名的 .txt 和 .xlsx 都支持, 优先用存在的那个
def _pick(stem):
    for ext in (".txt", ".tsv", ".csv", ".xlsx"):
        cand = os.path.join(DATA_DIR, stem + ext)
        if os.path.isfile(cand):
            return cand
    return os.path.join(DATA_DIR, stem + ".xlsx")


EXPR_FILE = _pick("expression_matrix")
MARKER_FILE = _pick("marker_genes")
if len(sys.argv) >= 3:                            # 也可用命令行参数覆盖
    EXPR_FILE, MARKER_FILE = sys.argv[1], sys.argv[2]
OUT_STEM = os.path.join(OUT_DIR, "Fig_coexpression_heatmap")

# ---------------- 参数 ----------------
TOPN = 20              # 每个 marker 取 |r| 最大的前 N 个基因
CORR_METHOD = "pearson"  # 在 log2(CPM+1) 上计算
GENE_LABEL_SIZE = 5.5  # 基因名过多, 单独设小字号; 其余文字为五号 10.5

# ---------------- 字体 ----------------
plt.rcParams["font.family"] = "serif"
plt.rcParams["font.serif"] = ["Times New Roman", "SimSun", "Liberation Serif"]
plt.rcParams["font.sans-serif"] = ["SimSun"]
plt.rcParams["mathtext.fontset"] = "stix"
plt.rcParams["font.weight"] = "bold"
plt.rcParams["axes.labelweight"] = "bold"
plt.rcParams["axes.titleweight"] = "bold"
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["font.size"] = 10.5
plt.rcParams["axes.labelsize"] = 10.5
plt.rcParams["xtick.labelsize"] = 10.5
plt.rcParams["ytick.labelsize"] = 10.5
plt.rcParams["pdf.fonttype"] = 42
plt.rcParams["ps.fonttype"] = 42
plt.rcParams["svg.fonttype"] = "none"


def _rows_from_file(path):
    """按扩展名读取 xlsx 或 文本表格 (制表符/逗号分隔), 统一返回二维列表"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xltx"):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    with open(path, "r", encoding="utf-8-sig") as fh:
        lines = [ln.rstrip("\r\n") for ln in fh if ln.strip()]
    sep = "\t" if "\t" in lines[0] else ("," if "," in lines[0] else None)
    return [ln.split(sep) if sep else [ln] for ln in lines]


def read_matrix(path):
    """读取表达矩阵, 返回 (样本名, 基因名列表, 计数矩阵)"""
    rows = _rows_from_file(path)
    header = [h for h in rows[0] if h not in (None, "")]
    samples = [str(h) for h in header[1:]]
    genes, values = [], []
    for r in rows[1:]:
        if r[0] in (None, ""):
            continue
        genes.append(str(r[0]).strip())
        values.append([float(r[i + 1]) for i in range(len(samples))])
    return samples, genes, np.array(values, dtype=float)


def read_marker_ids(path):
    """读取 marker 基因名。支持两种格式:
       1) 只有基因名的单列文件 (可有可无表头)
       2) 与表达矩阵同结构的文件, 取第一列"""
    rows = _rows_from_file(path)
    ids = []
    for i, r in enumerate(rows):
        first = str(r[0]).strip() if r and r[0] is not None else ""
        if not first:
            continue
        if i == 0 and (first.startswith("#") or first.lower() in ("id", "gene", "geneid")):
            continue                      # 跳过表头
        ids.append(first)
    return ids


def pearson_vs_matrix(v, M):
    """向量 v 与矩阵 M 每一行的 Pearson r"""
    vc = v - v.mean()
    Mc = M - M.mean(axis=1, keepdims=True)
    denom = np.sqrt((Mc ** 2).sum(axis=1)) * np.sqrt((vc ** 2).sum())
    denom[denom == 0] = np.nan
    return (Mc @ vc) / denom


def order_by_cluster(M):
    """对列做层次聚类排序; 若无 scipy 则退化为按最大相关的 marker 分组排序"""
    try:
        from scipy.cluster.hierarchy import linkage, leaves_list
        from scipy.spatial.distance import pdist
        d = pdist(M.T, metric="correlation")
        d = np.nan_to_num(d, nan=1.0)
        return leaves_list(linkage(d, method="average", optimal_ordering=True))
    except Exception:
        key = np.argmax(np.abs(M), axis=0)
        val = M[key, np.arange(M.shape[1])]
        return np.lexsort((-val, key))


# ---------------- 1. 读入并标准化 ----------------
samples, genes, counts = read_matrix(EXPR_FILE)
marker_ids = read_marker_ids(MARKER_FILE)

cpm = counts / counts.sum(axis=0, keepdims=True) * 1e6
logcpm = np.log2(cpm + 1.0)
gene_index = {g: i for i, g in enumerate(genes)}

missing = [g for g in marker_ids if g not in gene_index]
if missing:
    raise ValueError("marker 基因不在表达矩阵中: {}".format(missing))

other_genes = [g for g in genes if g not in set(marker_ids)]
other_mat = logcpm[[gene_index[g] for g in other_genes]]

# ---------------- 2. 相关系数矩阵 (6 x 其余基因) ----------------
R = np.vstack([pearson_vs_matrix(logcpm[gene_index[m]], other_mat) for m in marker_ids])

# ---------------- 3. 每个 marker 取 topN, 求并集 ----------------
selected = []
for i in range(len(marker_ids)):
    order = np.argsort(-np.abs(np.nan_to_num(R[i])))[:TOPN]
    for j in order:
        if other_genes[j] not in selected:
            selected.append(other_genes[j])
sel_idx = [other_genes.index(g) for g in selected]
Rsel = R[:, sel_idx]
col_order = order_by_cluster(Rsel)
Rsel = Rsel[:, col_order]
sel_names = [selected[j] for j in col_order]

# ---------------- 4. 绘图 ----------------
# 配色: NPG 蓝白红, 蓝为负相关, 红为正相关, 白为 0
# Pearson r 本身的取值范围是 [-1, 1], 故色标固定为 -1 到 1, 不随数据缩放
cmap = LinearSegmentedColormap.from_list(
    "npg_blue_white_red", ["#3C5488", "#FFFFFF", "#E64B35"])
norm = TwoSlopeNorm(vmin=-1.0, vcenter=0.0, vmax=1.0)

n_col = Rsel.shape[1]
n_row = Rsel.shape[0]
GENE_LABEL = "Co-expressed genes (top {} per marker, n = {})".format(TOPN, n_col)
MARKER_LABEL = "Marker genes"


def draw_heatmap(vertical):
    """vertical=False: marker 作行, 基因作列 (横版)
       vertical=True : 基因作行, marker 作列 (竖版)"""
    M = Rsel.T if vertical else Rsel
    nr, nc = M.shape
    if vertical:
        fig_w = n_row * 0.28 + 2.4
        fig_h = max(4.0, n_col * 0.105 + 1.4)
    else:
        fig_w = max(6.0, n_col * 0.105 + 1.8)
        fig_h = n_row * 0.28 + 1.9
    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=300)
    im = ax.imshow(M, cmap=cmap, norm=norm, aspect="auto", interpolation="nearest")

    if vertical:
        ax.set_xticks(np.arange(nc))
        ax.set_xticklabels(marker_ids, rotation=90, fontsize=8, fontweight="bold")
        ax.set_yticks(np.arange(nr))
        ax.set_yticklabels(sel_names, fontsize=GENE_LABEL_SIZE, fontweight="bold")
        ax.set_xlabel(MARKER_LABEL, fontsize=10.5, labelpad=6, fontweight="bold")
        ax.set_ylabel(GENE_LABEL, fontsize=10.5, fontweight="bold")
    else:
        ax.set_xticks(np.arange(nc))
        ax.set_xticklabels(sel_names, rotation=90, fontsize=GENE_LABEL_SIZE, fontweight="bold")
        ax.set_yticks(np.arange(nr))
        ax.set_yticklabels(marker_ids, fontsize=8, fontweight="bold")
        ax.set_xlabel(GENE_LABEL, fontsize=10.5, labelpad=6, fontweight="bold")
        ax.set_ylabel(MARKER_LABEL, fontsize=10.5, fontweight="bold")

    ax.set_xticks(np.arange(-0.5, nc, 1), minor=True)
    ax.set_yticks(np.arange(-0.5, nr, 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=0.4)
    ax.tick_params(which="minor", length=0)
    ax.tick_params(which="major", length=2.0, width=0.8)
    for side in ("left", "bottom", "right", "top"):
        ax.spines[side].set_linewidth(0.8)

    frac = 0.05 if vertical else 0.02
    cbar = fig.colorbar(im, ax=ax, fraction=frac, pad=0.02,
                        ticks=[-1.0, -0.5, 0.0, 0.5, 1.0])
    cbar.set_label("Pearson r", fontsize=10.5, fontweight="bold")
    cbar.ax.tick_params(labelsize=9)
    for lab in cbar.ax.get_yticklabels():
        lab.set_fontweight("bold")
    cbar.outline.set_linewidth(0.8)

    fig.tight_layout()
    stem = OUT_STEM + ("_vertical" if vertical else "_horizontal")
    for ext in ("pdf", "svg", "png"):
        path = "{}.{}".format(stem, ext)
        fig.savefig(path, format=ext, bbox_inches="tight", pad_inches=0.05)
        print("saved:", path)
    plt.close(fig)


draw_heatmap(vertical=False)
draw_heatmap(vertical=True)

# ---------------- 5. 导出完整相关系数表 ----------------
# 相关系数 -> t 检验 P 值 (df = n - 2), 再做 BH 多重检验校正
n_sample = logcpm.shape[1]
df_t = n_sample - 2
with np.errstate(divide="ignore", invalid="ignore"):
    tstat = R * np.sqrt(df_t) / np.sqrt(np.maximum(1.0 - R ** 2, 1e-12))
try:
    from scipy import stats as _stats
    P = 2.0 * _stats.t.sf(np.abs(tstat), df_t)
except Exception:                      # 无 scipy 时用正态近似, 仅作参考
    P = 2.0 * 0.5 * (1.0 - np.math.erf(np.abs(tstat) / np.sqrt(2.0)))
flat = P.ravel()
order = np.argsort(flat)
mtest = flat.size
Q = np.empty(mtest)
running = 1.0
for rank, pos in enumerate(order[::-1]):
    running = min(running, flat[pos] * mtest / (mtest - rank))
    Q[pos] = running
Q = Q.reshape(P.shape)

csv_path = OUT_STEM + "_correlation_full.csv"
with open(csv_path, "w", encoding="utf-8-sig") as fh:
    head = ["gene"]
    for m in marker_ids:
        head += ["r_" + m, "P_" + m, "BH_q_" + m]
    head.append("in_heatmap")
    fh.write(",".join(head) + "\n")
    for j, g in enumerate(other_genes):
        cells = [g]
        for i in range(n_row):
            cells += ["{:.4f}".format(R[i, j]), "{:.4g}".format(P[i, j]),
                      "{:.4g}".format(Q[i, j])]
        cells.append("yes" if g in selected else "no")
        fh.write(",".join(cells) + "\n")
print("saved:", csv_path)

# ---------------- 6. 关键信息输出 ----------------
print("\nsamples:", samples)
print("library size (raw counts):", counts.sum(axis=0).astype(int).tolist())
print("genes in matrix:", len(genes), " markers:", len(marker_ids),
      " others:", len(other_genes))
print("genes kept in heatmap:", n_col)
print("n = {} samples -> |r| > 0.811 corresponds to P < 0.05 (two-sided)".format(n_sample))
print("pairs with P < 0.05: {} / {}".format(int((P < 0.05).sum()), P.size))
print("pairs with BH q < 0.05: {} / {}".format(int((Q < 0.05).sum()), Q.size))
print("\ntop 3 co-expressed genes per marker:")
for i, m in enumerate(marker_ids):
    order = np.argsort(-np.abs(np.nan_to_num(R[i])))[:3]
    txt = ", ".join("{} (r = {:.3f})".format(other_genes[j], R[i, j]) for j in order)
    print("  {}: {}".format(m, txt))
