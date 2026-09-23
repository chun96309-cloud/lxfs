# -*- coding: utf-8 -*-
"""
每个基因一个小框的 FPKM 箱线图 (箱体 + 散点 + 显著性标记), 横向排列。

输入: 表格第一列为基因 ID, 样本列名形如 N1_FPKM / P1_FPKM (前缀字母为分组),
      若含 P 值列 (列名含 "Pvalue"), 显著性星号直接用该列; 否则用 Welch t 检验 (log2(FPKM+1))。
输出: <WORK_ROOT>\\当天日期\\Fig_fpkm_boxplot.pdf / .svg / .png
用法: python fpkm_boxplot.py [表格文件]   不带参数时读 data\\fpkm_boxplot.xlsx (或 .txt)
"""

import os
import re
import sys
import datetime
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch

# ---------------- 工作目录 ----------------
WORK_ROOT = r"C:\Users\23027\Desktop\11\画各种图"
os.makedirs(WORK_ROOT, exist_ok=True)
os.chdir(WORK_ROOT)
DATA_DIR = os.path.join(WORK_ROOT, "data")
os.makedirs(DATA_DIR, exist_ok=True)
OUT_DIR = os.path.join(WORK_ROOT, datetime.date.today().strftime("%Y-%m-%d"))
os.makedirs(OUT_DIR, exist_ok=True)


def _pick(stem):
    for ext in (".xlsx", ".txt", ".tsv", ".csv"):
        cand = os.path.join(DATA_DIR, stem + ext)
        if os.path.isfile(cand):
            return cand
    return os.path.join(DATA_DIR, stem + ".xlsx")


IN_FILE = sys.argv[1] if len(sys.argv) >= 2 else _pick("fpkm_boxplot")
OUT_STEM = os.path.join(OUT_DIR, "Fig_fpkm_boxplot")

# ---------------- 参数 ----------------
NCOL = None                 # 每行放几个小框; None = 全部一行
PANEL_W, PANEL_H = 1.45, 2.4   # 单个小框尺寸 (英寸)
GROUP_COLORS = {"N": "#7EA1C4", "P": "#6A6BB0"}   # 参考图左(浅蓝)与中(蓝紫)两色
DEFAULT_COLORS = ["#7EA1C4", "#6A6BB0", "#6B5A8E", "#91ABD2"]
Y_LABEL = "FPKM"
POINT_SIZE = 16
JITTER = 0.10
SEED = 0

# ---------------- 字体 ----------------
plt.rcParams["font.family"] = "serif"
plt.rcParams["font.serif"] = ["Times New Roman", "SimSun", "Liberation Serif"]
plt.rcParams["font.sans-serif"] = ["SimSun"]
plt.rcParams["mathtext.fontset"] = "stix"
plt.rcParams["font.weight"] = "bold"
plt.rcParams["axes.labelweight"] = "bold"
plt.rcParams["axes.titleweight"] = "bold"
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["font.size"] = 10.5
plt.rcParams["pdf.fonttype"] = 42
plt.rcParams["ps.fonttype"] = 42
plt.rcParams["svg.fonttype"] = "none"

_META = {"pdf": {"Creator": None, "Producer": None},
         "svg": {"Creator": None, "Date": None},
         "png": {"Software": None}}


def _scrub_svg(path):
    with open(path, "r", encoding="utf-8") as fh:
        txt = fh.read()
    txt = txt.replace('id="matplotlib.', 'id="').replace("#matplotlib.", "#")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(txt)


# ---------------- 读取 ----------------
def _rows_from_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xltx"):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    with open(path, "r", encoding="utf-8-sig") as fh:
        lines = [ln.rstrip("\r\n") for ln in fh if ln.strip()]
    sep = "\t" if "\t" in lines[0] else ","
    return [ln.split(sep) for ln in lines]


rows = _rows_from_file(IN_FILE)
header = [str(h).strip() if h is not None else "" for h in rows[0]]
sample_cols = [i for i, h in enumerate(header) if h.upper().endswith("_FPKM")]
if not sample_cols:
    raise SystemExit("没有找到以 _FPKM 结尾的样本列")
p_cols = [i for i, h in enumerate(header)
          if "pvalue" in h.lower() and "regulated" not in h.lower()]
p_col = p_cols[0] if p_cols else None


def group_of(colname):
    m = re.match(r"([A-Za-z]+)", colname)
    return m.group(1) if m else colname


groups = []
for i in sample_cols:
    g = group_of(header[i])
    if g not in groups:
        groups.append(g)
colors = {}
for k, g in enumerate(groups):
    colors[g] = GROUP_COLORS.get(g, DEFAULT_COLORS[k % len(DEFAULT_COLORS)])

genes, data, pvals = [], [], []
for r in rows[1:]:
    if r is None or r[0] in (None, ""):
        continue
    genes.append(str(r[0]).strip())
    d = {g: [] for g in groups}
    for i in sample_cols:
        v = r[i]
        if v is None or v == "":
            continue
        d[group_of(header[i])].append(float(v))
    data.append(d)
    pvals.append(float(r[p_col]) if p_col is not None and r[p_col] not in (None, "") else None)

# ---------------- 显著性 ----------------
from scipy import stats as _stats


def stars(p):
    if p is None or np.isnan(p):
        return "ns"
    if p < 1e-4:
        return "****"
    if p < 1e-3:
        return "***"
    if p < 1e-2:
        return "**"
    if p < 0.05:
        return "*"
    return "ns"


sig_source = "DESeq2 P value (from input table)" if p_col is not None else "Welch t-test on log2(FPKM+1)"
if p_col is None:
    pvals = []
    for d in data:
        a = np.log2(np.array(d[groups[0]]) + 1)
        b = np.log2(np.array(d[groups[1]]) + 1)
        pvals.append(float(_stats.ttest_ind(a, b, equal_var=False).pvalue))

# ---------------- 绘图 ----------------
n = len(genes)
ncol = n if NCOL is None else int(NCOL)
nrow = int(np.ceil(n / ncol))
fig, axes = plt.subplots(nrow, ncol, figsize=(PANEL_W * ncol + 0.9, PANEL_H * nrow + 0.5),
                         dpi=300, squeeze=False)
rng = np.random.default_rng(SEED)

for k, gene in enumerate(genes):
    ax = axes[k // ncol][k % ncol]
    d = data[k]
    vals = [np.array(d[g]) for g in groups]
    pos = np.arange(len(groups))
    bp = ax.boxplot(vals, positions=pos, widths=0.55, patch_artist=True, showfliers=False,
                    medianprops=dict(color="none"), whiskerprops=dict(linewidth=1.2),
                    capprops=dict(linewidth=1.2), boxprops=dict(linewidth=1.2), zorder=2)
    for i, g in enumerate(groups):
        c = colors[g]
        bp["boxes"][i].set_facecolor(c)
        bp["boxes"][i].set_alpha(0.45)
        bp["boxes"][i].set_edgecolor(c)
        for part in ("whiskers", "caps"):
            for art in bp[part][2 * i:2 * i + 2]:
                art.set_color(c)
        ax.hlines(np.median(vals[i]), pos[i] - 0.275, pos[i] + 0.275, color=c, linewidth=2.0, zorder=3)
        x = pos[i] + rng.uniform(-JITTER, JITTER, size=len(vals[i]))
        ax.scatter(x, vals[i], s=POINT_SIZE, color=c, alpha=0.8, edgecolor="none", zorder=4)

    # 显著性横线
    ymax = max(v.max() for v in vals)
    ymin = min(v.min() for v in vals)
    span = max(ymax - ymin, 1e-9)
    y0 = ymax + span * 0.10
    h = span * 0.05
    ax.plot([pos[0], pos[0], pos[-1], pos[-1]], [y0, y0 + h, y0 + h, y0], color="black", linewidth=1.0)
    ax.text((pos[0] + pos[-1]) / 2.0, y0 + h * 1.15, stars(pvals[k]), ha="center", va="bottom",
            fontsize=10.5, fontweight="bold")
    ax.set_ylim(ymin - span * 0.08, y0 + h + span * 0.30)

    ax.set_title(gene, fontsize=9, fontweight="bold", pad=4)
    ax.set_xticks(pos)
    ax.set_xticklabels(groups, fontsize=10.5, fontweight="bold")
    ax.set_xlim(-0.6, len(groups) - 0.4)
    ax.tick_params(axis="y", labelsize=8, width=0.8, length=3)
    ax.tick_params(axis="x", width=0.8, length=3)
    for lab in ax.get_yticklabels():
        lab.set_fontweight("bold")
    ax.grid(True, axis="y", color="#E5E5E5", linewidth=0.6, zorder=0)
    ax.set_axisbelow(True)
    for side in ("left", "bottom", "right", "top"):
        ax.spines[side].set_linewidth(0.9)
    if k % ncol == 0:
        ax.set_ylabel(Y_LABEL, fontsize=10.5, fontweight="bold")

# 多余的格子隐藏
for k in range(n, nrow * ncol):
    axes[k // ncol][k % ncol].axis("off")

handles = [Patch(facecolor=colors[g], edgecolor=colors[g], alpha=0.6, label=g) for g in groups]
fig.legend(handles=handles, loc="center right", frameon=False, fontsize=10.5,
           bbox_to_anchor=(1.0, 0.5), handlelength=1.2)
fig.tight_layout(rect=(0, 0, 0.955, 1), w_pad=0.6)

for ext in ("pdf", "svg", "png"):
    path = "{}.{}".format(OUT_STEM, ext)
    fig.savefig(path, format=ext, bbox_inches="tight", pad_inches=0.05, metadata=_META[ext])
    if ext == "svg":
        _scrub_svg(path)
    print("saved:", path)
plt.close(fig)

# ---------------- 关键信息 ----------------
print("\ninput:", IN_FILE)
print("groups:", {g: sum(1 for i in sample_cols if group_of(header[i]) == g) for g in groups})
print("significance source:", sig_source)
for gene, d, p in zip(genes, data, pvals):
    print("  {}: {} | P = {:.3g} -> {}".format(
        gene, " ".join("{} mean {:.2f}".format(g, np.mean(d[g])) for g in groups), p, stars(p)))
